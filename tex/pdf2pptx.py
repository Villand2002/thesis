"""Beamer PDF -> PPTX that keeps the Beamer look.

Each slide = the PDF page rendered with its text removed (title bar, blocks,
ball bullets, table rules, footline stay pixel-identical) + the text re-placed
as editable PowerPoint text at the same positions, with real super/subscripts.
Speaker notes come from the \\note{} blocks of slides.tex.
"""
import re
import sys
import pymupdf
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE

PDF, TEX, OUT, TMP = sys.argv[1:5]

# ---------- speaker notes ----------
def latex_to_plain(s):
    s = re.sub(r"\\emph\{([^}]*)\}", r"\1", s)
    s = s.replace("``", "\u201c").replace("''", "\u201d").replace("---", "\u2014").replace("--", "\u2013")
    s = s.replace('\\"u', "ü").replace('\\"o', "ö")
    s = re.sub(r"\$([^$]*)\$", lambda m: m.group(1).replace("\\delta", "δ").replace("\\sigma", "σ")
               .replace("i_", "i").replace("t_", "t"), s)
    return re.sub(r"\s+", " ", s).strip()

notes = [latex_to_plain(m) for m in re.findall(r"\\note\{([\s\S]*?)\}\n", open(TEX).read())]

# ---------- fonts ----------
SUP = "/System/Library/Fonts/Supplemental/"
FONTFILES = {
    ("Arial", False, False): SUP + "Arial.ttf",
    ("Arial", True, False): SUP + "Arial Bold.ttf",
    ("Arial", False, True): SUP + "Arial Italic.ttf",
    ("Arial Unicode MS", False, False): SUP + "Arial Unicode.ttf",
    ("Cambria Math", False, False): "/Applications/Microsoft PowerPoint.app/Contents/Resources/DFonts/Cambria.ttc",
}
_fc = {}
def measure(text, face, bold, italic, size):
    key = (face, bold, italic) if (face, bold, italic) in FONTFILES else (face, False, False)
    f = _fc.setdefault((key, round(size * 4)), ImageFont.truetype(FONTFILES[key], max(1, round(size * 4)), index=1 if key[0] == "Cambria Math" else 0))
    return f.getlength(text) / 4.0

arial_cmap = set(pymupdf.Font(fontfile=SUP + "Arial.ttf").valid_codepoints())

LIG = {"ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl"}
def clean(t):
    for a, b in LIG.items():
        t = t.replace(a, b)
    return t.replace("⇐⇒", "\u21d4").replace("¨u", "ü").replace("¨o", "ö").replace("µ", "μ")

SCRIPT_EXC = {"B": "\u212c", "E": "\u2130", "F": "\u2131", "H": "\u210b", "I": "\u2110",
              "L": "\u2112", "M": "\u2133", "R": "\u211b"}
def calligraphic(text):
    return "".join(SCRIPT_EXC.get(c, chr(0x1D49C + ord(c) - 65)) if "A" <= c <= "Z" else c for c in text)

def style(fontname, text):
    """Map a Computer Modern font to (face, bold, italic)."""
    fn = fontname.upper()
    if "CMSY" in fn or "CMEX" in fn:
        if any(ord(c) > 0xFFFF or 0x2100 <= ord(c) <= 0x214F and c not in "\u2122" for c in text):
            face = "Cambria Math"
        else:
            face = "Arial" if all(ord(c) in arial_cmap for c in text if not c.isspace()) else "Arial Unicode MS"
        return face, False, False
    if "CMSSBX" in fn or "BOLD" in fn:
        return "Arial", True, False
    if "CMSSI" in fn or "CMMI" in fn or "OBLIQUE" in fn or "ITALIC" in fn:
        return "Arial", False, True
    return "Arial", False, False

# ---------- geometry ----------
src = pymupdf.open(PDF)
PW, PH = src[0].rect.width, src[0].rect.height
SLIDE_W = Emu(9144000)                      # 10 in
SLIDE_H = Emu(round(9144000 * PH / PW))     # 7.5 in for 4:3
F = 720.0 / PW                              # PDF pt -> slide pt
ASC = 0.967                                 # Arial: first baseline = top + ASC*size (measured in PowerPoint, zero inset)
SCRIPT_SHRINK = 1.5                         # PowerPoint renders a super/subscript run at ~2/3 of its size

def keep_in_background(s):
    x0, y0, x1, y1 = s["bbox"]
    if s["origin"][1] > PH - 10:            # footline
        return True
    if s["color"] == 0xFFFFFF and s["size"] < 6.5 and x1 < 30:   # numbers inside enumerate balls
        return True
    return False

prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
blank = prs.slide_layouts[6]

for pno in range(len(src)):
    page = src[pno]
    spans = []
    for b in page.get_text("rawdict")["blocks"]:
        for l in b.get("lines", []):
            for s in l["spans"]:
                txt, prev = "", None
                for ch in s["chars"]:
                    if prev is not None and ch["c"] != " " and not txt.endswith(" ") \
                            and ch["bbox"][0] - prev > 0.2 * s["size"]:
                        txt += " "
                    txt += ch["c"]
                    prev = ch["bbox"][2]
                s["text"] = txt
                if not s["text"].strip() or keep_in_background(s):
                    continue
                spans.append(s)

    # background: same page, text removed
    bg = pymupdf.open()
    bg.insert_pdf(src, from_page=pno, to_page=pno)
    bp = bg[0]
    for s in spans:
        bp.add_redact_annot(pymupdf.Rect(s["bbox"]), fill=False)
    bp.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
    img = f"{TMP}/bg-{pno + 1:02d}.png"
    bp.get_pixmap(dpi=400).save(img)

    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(img, 0, 0, SLIDE_W, SLIDE_H)

    # ---- group spans into lines; smaller spans near a baseline become scripts ----
    items = [dict(text=calligraphic(clean(s["text"])) if "CMSY" in s["font"] else clean(s["text"]), font=s["font"], size=s["size"], x0=s["bbox"][0], x1=s["bbox"][2], y0=s["bbox"][1], y1=s["bbox"][3],
                  b=s["origin"][1], color=s["color"]) for s in spans]
    items.sort(key=lambda it: ("CMEX" in it["font"], -round(it["size"], 1), it["b"], it["x0"]))
    lines = []
    for it in items:
        home = None
        if "CMEX" in it["font"]:
            cands = [L for L in lines if it["y0"] - 1 <= L["b"] <= it["y1"] + 10 and L["b"] >= it["y1"] - 2]
            if cands:
                home = min(cands, key=lambda L: L["b"] - it["y1"]); it["raise"] = 0.0
            if home is not None:
                home["members"].append(it)
                home["x0"] = min(home["x0"], it["x0"]); home["x1"] = max(home["x1"], it["x1"])
                continue
        for L in lines:
            if abs(L["b"] - it["b"]) < 0.15 * L["size"] and it["size"] >= 0.85 * L["size"]:
                home = L; it["raise"] = 0.0; break
        if home is None:
            for L in lines:
                if (it["size"] < 0.85 * L["size"] and abs(L["b"] - it["b"]) < 0.75 * L["size"]
                        and L["x0"] - L["size"] <= it["x0"] <= L["x1"] + L["size"]):
                    home = L; it["raise"] = L["b"] - it["b"]; break
        if home is None:
            home = dict(b=it["b"], size=it["size"], x0=it["x0"], x1=it["x1"], members=[])
            it["raise"] = 0.0
            lines.append(home)
        home["members"].append(it)
        home["x0"] = min(home["x0"], it["x0"]); home["x1"] = max(home["x1"], it["x1"])

    for L in lines:
        mem = sorted(L["members"], key=lambda m: m["x0"])
        i = 0
        while i < len(mem):
            if mem[i]["raise"] != 0.0:
                j = i
                while j < len(mem) and mem[j]["raise"] != 0.0:
                    j += 1
                grp = mem[i:j]
                if any(g["raise"] < 0 for g in grp) and any(g["raise"] > 0 for g in grp):
                    mem[i:j] = [g for g in grp if g["raise"] < 0] + [g for g in grp if g["raise"] > 0]
                i = j
            else:
                i += 1
        # split into segments at large horizontal gaps (table cells, columns)
        segs, cur = [], [mem[0]]
        for m in mem[1:]:
            if m["x0"] - max(c["x1"] for c in cur) > 0.9 * L["size"]:
                segs.append(cur); cur = [m]
            else:
                cur.append(m)
        segs.append(cur)

        for seg in segs:
            main = L["size"] * F
            runs, natural, nchar, prev_x1 = [], 0.0, 0, None
            for m in seg:
                face, bold, italic = style(m["font"], m["text"])
                t = m["text"].strip()
                if prev_x1 is not None and m["x0"] - prev_x1 > (0.3 if runs and ord(runs[-1][0][-1:] or " ") > 0xFFFF else 0.18) * L["size"] and not m["text"].startswith(" "):
                    t = " " + t
                elif m["text"].startswith(" ") and prev_x1 is not None:
                    t = " " + t
                prev_x1 = m["x1"]
                shown = m["size"] * F                      # size as it should appear
                if m["raise"] != 0.0:
                    spec = shown * SCRIPT_SHRINK
                    baseline = round(m["raise"] * F / spec * 100000)
                else:
                    spec, baseline = shown, None
                runs.append((t, face, bold, italic, spec, baseline, m["color"]))
                natural += measure(t, face, bold, italic, shown)
                nchar += len(t)
            target = (seg[-1]["x1"] - seg[0]["x0"]) * F
            spc = max(-2.0, min(2.0, (target - natural) / max(1, nchar)))

            left = seg[0]["x0"] * F
            top = L["b"] * F - ASC * main
            tb = slide.shapes.add_textbox(Pt(left), Pt(top), Pt(target + 12), Pt(main * 1.4))
            tf = tb.text_frame
            tf.word_wrap = False
            tf.auto_size = MSO_AUTO_SIZE.NONE
            tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            for t, face, bold, italic, spec, baseline, color in runs:
                r = p.add_run()
                r.text = t
                r.font.name = face
                r.font.size = Pt(round(spec * 2) / 2)
                r.font.bold = bold
                r.font.italic = italic
                r.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
                rpr = r.font._rPr
                rpr.set("spc", str(round(spc * 100)))
                if baseline is not None:
                    rpr.set("baseline", str(baseline))

    if pno < len(notes):
        slide.notes_slide.notes_text_frame.text = notes[pno]

prs.save(OUT)
print("wrote", OUT, len(prs.slides), "slides,", len(notes), "notes")
